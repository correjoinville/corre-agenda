import csv, io, json, math, os, re, sqlite3, statistics, unicodedata, uuid
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from functools import wraps
from pathlib import Path
from urllib.parse import urlparse
import click
from flask import Flask, abort, current_app, flash, g, redirect, render_template, request, session, url_for
from flask_wtf.csrf import CSRFProtect
from openpyxl import load_workbook
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

MONTHS_PT=["","Janeiro","Fevereiro","Março","Abril","Maio","Junho","Julho","Agosto","Setembro","Outubro","Novembro","Dezembro"]
RESULT_FIELDS=[("bib","Número do atleta"),("name","Nome"),("sex","Sexo"),("age","Idade"),("category","Categoria"),("distance","Distância"),("gross_time","Tempo bruto"),("net_time","Tempo líquido"),("overall_place","Colocação geral"),("sex_place","Colocação por sexo"),("category_place","Colocação na categoria"),("status","Status")]
ALIASES={"bib":["numero","número","peito","bib"],"name":["nome","atleta","participante"],"sex":["sexo","genero","gênero"],"age":["idade"],"category":["categoria","faixa"],"distance":["distancia","distância","prova","percurso"],"gross_time":["tempo bruto","tempo oficial","bruto"],"net_time":["tempo liquido","tempo líquido","tempo","liquido","líquido"],"overall_place":["colocacao geral","colocação geral","geral"],"sex_place":["colocacao sexo","colocação sexo"],"category_place":["colocacao categoria","colocação categoria"],"status":["status","situacao","situação"]}

def create_app(test_config=None):
    app=Flask(__name__)
    app.config.from_mapping(SECRET_KEY=os.environ.get("SECRET_KEY"),DATABASE=os.environ.get("DATABASE_PATH",os.path.join(app.root_path,"database.db")),IMPORT_FOLDER=os.path.join(app.instance_path,"imports"),SESSION_COOKIE_HTTPONLY=True,SESSION_COOKIE_SAMESITE="Lax",SESSION_COOKIE_SECURE=os.environ.get("FLASK_ENV")=="production",PERMANENT_SESSION_LIFETIME=timedelta(hours=8),MAX_CONTENT_LENGTH=10*1024*1024)
    if test_config: app.config.update(test_config)
    if not app.config.get("SECRET_KEY"): raise RuntimeError("Defina a variável de ambiente SECRET_KEY antes de iniciar o site.")
    Path(app.config["IMPORT_FOLDER"]).mkdir(parents=True,exist_ok=True); CSRFProtect(app); app.teardown_appcontext(close_db); app.cli.add_command(init_db_command)
    with app.app_context(): ensure_schema()
    app.add_template_filter(lambda v:datetime.strptime(v,"%Y-%m-%d").strftime("%d/%m/%Y"),"date_br")
    app.add_template_filter(format_seconds,"time_fmt")

    @app.get("/")
    def index():
        today=get_today(app); cutoff=(today-timedelta(days=15)).isoformat()
        events=get_db().execute("""SELECT e.*,EXISTS(SELECT 1 FROM results r WHERE r.event_id=e.id) has_results FROM events e WHERE status='publicado' AND event_date>=? ORDER BY event_date,name COLLATE NOCASE""",(cutoff,)).fetchall(); groups=[]
        for event in events:
            d=date.fromisoformat(event["event_date"]); key=(d.year,d.month)
            if not groups or groups[-1]["key"]!=key: groups.append({"key":key,"title":f"{MONTHS_PT[d.month]} {d.year}","events":[]})
            groups[-1]["events"].append(event)
        return render_template("index.html",groups=groups,today=today,months=MONTHS_PT)

    @app.route("/login",methods=("GET","POST"))
    def login():
        if session.get("admin_id"): return redirect(url_for("admin"))
        if request.method=="POST":
            u=request.form.get("username","").strip(); user=get_db().execute("SELECT * FROM admins WHERE username=?",(u,)).fetchone()
            if user and check_password_hash(user["password_hash"],request.form.get("password","")):
                session.clear(); session["admin_id"]=user["id"]; session.permanent=True; return redirect(url_for("admin"))
            flash("Usuário ou senha inválidos.","error")
        return render_template("login.html")

    @app.post("/logout")
    @login_required
    def logout(): session.clear(); flash("Você saiu da área administrativa.","success"); return redirect(url_for("login"))

    @app.get("/admin")
    @login_required
    def admin():
        events=get_db().execute("""SELECT e.*,COUNT(r.id) result_count,SUM(CASE WHEN r.status='FINISHED' THEN 1 ELSE 0 END) finisher_count FROM events e LEFT JOIN results r ON r.event_id=e.id GROUP BY e.id ORDER BY e.event_date DESC,e.name COLLATE NOCASE""").fetchall()
        return render_template("admin.html",events=events)

    @app.route("/admin/eventos/novo",methods=("GET","POST"))
    @login_required
    def event_create():
        if request.method=="POST":
            data,errors=validate_event_form(request.form)
            if not errors:
                now=datetime.now().isoformat(timespec="seconds"); get_db().execute("INSERT INTO events(event_date,name,city,state,registration_url,photos_url,status,slug,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?)",(*data,unique_slug(data[1]),now,now)); get_db().commit(); flash("Corrida adicionada com sucesso.","success"); return redirect(url_for("admin"))
            for error in errors: flash(error,"error")
        return render_template("event_form.html",event=None,form=request.form)

    @app.route("/admin/eventos/<int:event_id>/editar",methods=("GET","POST"))
    @login_required
    def event_edit(event_id):
        event=get_event_or_404(event_id)
        if request.method=="POST":
            data,errors=validate_event_form(request.form)
            if not errors:
                get_db().execute("UPDATE events SET event_date=?,name=?,city=?,state=?,registration_url=?,photos_url=?,status=?,updated_at=? WHERE id=?",(*data,datetime.now().isoformat(timespec="seconds"),event_id)); get_db().commit(); flash("Corrida atualizada com sucesso.","success"); return redirect(url_for("admin"))
            for error in errors: flash(error,"error")
        return render_template("event_form.html",event=event,form=request.form)

    @app.post("/admin/eventos/<int:event_id>/excluir")
    @login_required
    def event_delete(event_id):
        event=get_event_or_404(event_id); get_db().execute("DELETE FROM events WHERE id=?",(event_id,)); get_db().commit(); flash(f'“{event["name"]}” foi excluída.',"success"); return redirect(url_for("admin"))

    @app.route("/admin/eventos/<int:event_id>/resultados/importar",methods=("GET","POST"))
    @login_required
    def import_results(event_id):
        event=get_event_or_404(event_id)
        if request.method=="POST":
            upload=request.files.get("file")
            if not upload or not upload.filename: flash("Selecione um arquivo CSV ou XLSX.","error")
            else:
                try:
                    headers,rows=read_results_file(upload)
                    if not headers or not rows: raise ValueError("O arquivo não contém resultados.")
                    token=save_import({"event_id":event_id,"filename":secure_filename(upload.filename),"headers":headers,"rows":rows}); return redirect(url_for("map_results",token=token))
                except (ValueError,OSError) as exc: flash(str(exc),"error")
        return render_template("import_upload.html",event=event)

    @app.route("/admin/resultados/mapear/<token>",methods=("GET","POST"))
    @login_required
    def map_results(token):
        payload=load_import(token); event=get_event_or_404(payload["event_id"]); suggested=suggest_mapping(payload["headers"])
        if request.method=="POST":
            mapping={f:request.form.get(f,"") for f,_ in RESULT_FIELDS}
            if not mapping.get("name") or not(mapping.get("net_time") or mapping.get("gross_time")): flash("Mapeie ao menos Nome e um campo de Tempo.","error")
            else:
                normalized,counts=normalize_rows(payload["rows"],mapping); payload.update(mapping=mapping,normalized=normalized,counts=counts); write_import(token,payload)
                return render_template("import_confirm.html",event=event,token=token,counts=counts,preview=normalized[:10])
        return render_template("import_mapping.html",event=event,token=token,headers=payload["headers"],fields=RESULT_FIELDS,suggested=suggested,sample=payload["rows"][:5])

    @app.post("/admin/resultados/confirmar/<token>")
    @login_required
    def confirm_results(token):
        payload=load_import(token)
        if "normalized" not in payload: abort(400)
        db=get_db(); event_id=payload["event_id"]
        if request.form.get("replace")=="1": db.execute("DELETE FROM results WHERE event_id=?",(event_id,))
        distances={}
        for item in payload["normalized"]:
            row=dict(item); label=row.pop("distance_label") or "Geral"; km=parse_distance_km(label); key=(label,km)
            if key not in distances:
                found=db.execute("SELECT id FROM distances WHERE event_id=? AND label=?",(event_id,label)).fetchone()
                if found: distances[key]=found[0]
                else: distances[key]=db.execute("INSERT INTO distances(event_id,label,distance_km) VALUES(?,?,?)",(event_id,label,km)).lastrowid
            db.execute("""INSERT INTO results(event_id,distance_id,bib,name,sex,age,category,gross_seconds,net_seconds,overall_place,sex_place,category_place,status,raw_status,imported_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",(event_id,distances[key],row["bib"],row["name"],row["sex"],row["age"],row["category"],row["gross_seconds"],row["net_seconds"],row["overall_place"],row["sex_place"],row["category_place"],row["status"],row["raw_status"],datetime.now().isoformat(timespec="seconds")))
        db.commit(); delete_import(token); flash(f'{len(payload["normalized"])} registros importados.',"success"); return redirect(url_for("admin"))

    @app.get("/resultados/<slug>")
    def results_dashboard(slug):
        event=get_db().execute("SELECT * FROM events WHERE slug=? AND status='publicado'",(slug,)).fetchone()
        if not event: abort(404)
        distances=get_db().execute("SELECT d.*,COUNT(r.id) total FROM distances d JOIN results r ON r.distance_id=d.id WHERE d.event_id=? GROUP BY d.id ORDER BY d.distance_km,d.label",(event["id"],)).fetchall()
        if not distances: abort(404)
        chosen=request.args.get("distancia",type=int); selected=next((d for d in distances if d["id"]==chosen),distances[0]); rows=get_db().execute("SELECT * FROM results WHERE distance_id=?",(selected["id"],)).fetchall(); q=request.args.get("q","").strip(); athletes=[]
        if q: athletes=get_db().execute("SELECT id,bib,name,sex,category FROM results WHERE distance_id=? AND (name LIKE ? OR bib LIKE ?) ORDER BY name LIMIT 20",(selected["id"],f"%{q}%",f"%{q}%")).fetchall()
        return render_template("results_dashboard.html",event=event,distances=distances,selected=selected,stats=calculate_stats(rows,selected["distance_km"]),q=q,athletes=athletes)

    @app.get("/resultados/<slug>/atleta/<int:result_id>")
    def athlete_result(slug,result_id):
        row=get_db().execute("""SELECT r.*,d.label distance_label,d.distance_km,e.name event_name,e.slug,e.event_date,e.city,e.state FROM results r JOIN distances d ON d.id=r.distance_id JOIN events e ON e.id=r.event_id WHERE r.id=? AND e.slug=? AND e.status='publicado'""",(result_id,slug)).fetchone()
        if not row: abort(404)
        rows=get_db().execute("SELECT * FROM results WHERE distance_id=?",(row["distance_id"],)).fetchall(); return render_template("athlete_result.html",athlete=row,analysis=athlete_analysis(row,rows))
    return app

def get_today(app): return date.fromisoformat(app.config["TODAY_OVERRIDE"]) if app.config.get("TODAY_OVERRIDE") else date.today()
def get_db():
    if "db" not in g: g.db=sqlite3.connect(current_app.config["DATABASE"]); g.db.row_factory=sqlite3.Row; g.db.execute("PRAGMA foreign_keys=ON")
    return g.db
def close_db(_error=None):
    db=g.pop("db",None)
    if db is not None: db.close()
def login_required(view):
    @wraps(view)
    def wrapped(**kwargs): return redirect(url_for("login",next=request.path)) if not session.get("admin_id") else view(**kwargs)
    return wrapped
def safe_url(value):
    if not value:return ""
    try:
        p=urlparse(value); return value if p.scheme in("http","https") and p.netloc else None
    except ValueError:return None
def validate_event_form(form):
    v=[form.get(k,"").strip() for k in("event_date","name","city","state","registration_url","photos_url","status")]; v[3]=v[3].upper(); e=[]
    try: date.fromisoformat(v[0])
    except ValueError:e.append("Informe uma data válida.")
    if not 2<=len(v[1])<=120:e.append("O nome deve ter entre 2 e 120 caracteres.")
    if not 2<=len(v[2])<=80:e.append("A cidade deve ter entre 2 e 80 caracteres.")
    if len(v[3])!=2 or not v[3].isalpha():e.append("Use a sigla do estado com 2 letras.")
    if safe_url(v[4]) is None:e.append("O link de inscrição deve começar com http:// ou https://.")
    if safe_url(v[5]) is None:e.append("O link das fotos deve começar com http:// ou https://.")
    if v[6] not in("publicado","rascunho"):e.append("Selecione um status válido.")
    return tuple(v),e
def get_event_or_404(event_id):
    event=get_db().execute("SELECT * FROM events WHERE id=?",(event_id,)).fetchone()
    if event is None:abort(404)
    return event
def slugify(value): return re.sub(r"[^a-z0-9]+","-",unicodedata.normalize("NFKD",value).encode("ascii","ignore").decode().lower()).strip("-") or "corrida"
def unique_slug(name,event_id=None):
    base=slugify(name); candidate=base; n=2
    while get_db().execute("SELECT 1 FROM events WHERE slug=? AND id!=?",(candidate,event_id or -1)).fetchone():candidate=f"{base}-{n}";n+=1
    return candidate
def parse_time(value):
    if value is None:return None
    if isinstance(value,(int,float)) and not isinstance(value,bool):return round(value*86400) if 0<value<1 else (int(value) if value>0 else None)
    text=str(value).strip().replace(",",".")
    if not text or text.upper() in{"DNS","DNF","DQ","DSQ","DESC","DESCLASSIFICADO","-"}:return None
    try:
        parts=text.split(":")
        if len(parts)==2:return round(float(parts[0])*60+float(parts[1]))
        if len(parts)==3:return round(float(parts[0])*3600+float(parts[1])*60+float(parts[2]))
    except ValueError:pass
    return None
def format_seconds(seconds):
    if seconds is None:return "—"
    seconds=int(round(seconds));h,rem=divmod(seconds,3600);m,s=divmod(rem,60);return f"{h}:{m:02d}:{s:02d}" if h else f"{m}:{s:02d}"
def parse_int(value):
    match=re.search(r"\d+",str(value or "").replace(".",""));return int(match.group()) if match else None
def parse_distance_km(value):
    match=re.search(r"(\d+(?:[.,]\d+)?)",str(value or ""));return float(match.group(1).replace(",",".")) if match else None
def canonical_status(raw,time_seconds):
    text=str(raw or "").strip().upper()
    if "DNS" in text or "NÃO LARG" in text or "NAO LARG" in text:return "DNS"
    if "DNF" in text or "NÃO CONCL" in text or "NAO CONCL" in text:return "DNF"
    if any(x in text for x in("DSQ","DQ","DESC","DESCLASS")):return "DQ"
    return "FINISHED" if time_seconds else "ERROR"
def read_results_file(upload):
    ext=Path(upload.filename).suffix.lower()
    if ext==".csv":
        raw=upload.read();text=None
        for enc in("utf-8-sig","latin-1"):
            try:text=raw.decode(enc);break
            except UnicodeDecodeError:pass
        try:dialect=csv.Sniffer().sniff(text[:4096],delimiters=",;\t|")
        except csv.Error:dialect=csv.excel
        reader=csv.DictReader(io.StringIO(text),dialect=dialect);headers=[str(h).strip() for h in(reader.fieldnames or [])];rows=[{str(k).strip():("" if v is None else str(v).strip()) for k,v in row.items()} for row in reader]
    elif ext==".xlsx":
        wb=load_workbook(upload,read_only=True,data_only=True);it=wb.active.iter_rows(values_only=True)
        try:headers=[str(v).strip() if v is not None else "" for v in next(it)]
        except StopIteration:return [],[]
        rows=[{headers[i]:("" if v is None else str(v)) for i,v in enumerate(vals) if i<len(headers)} for vals in it]
    else:raise ValueError("Formato inválido. Envie um arquivo CSV ou XLSX.")
    headers=[h for h in headers if h];return headers,[r for r in rows if any(str(v).strip() for v in r.values())]
def suggest_mapping(headers):
    normalized={slugify(h).replace("-"," "):h for h in headers};result={}
    for field,aliases in ALIASES.items():
        for alias in aliases:
            key=slugify(alias).replace("-"," ")
            if key in normalized:result[field]=normalized[key];break
        if field not in result:
            for key,original in normalized.items():
                if any(slugify(a).replace("-"," ") in key for a in aliases):result[field]=original;break
    return result
def normalize_rows(rows,mapping):
    output=[];counts=Counter(total=len(rows))
    for source in rows:
        get=lambda f:source.get(mapping.get(f,""),"") if mapping.get(f) else "";gross,net=parse_time(get("gross_time")),parse_time(get("net_time"));chosen=net or gross;raw=get("status") or(get("net_time") if not chosen else "");status=canonical_status(raw,chosen)
        if not str(get("name")).strip() or status=="ERROR":counts["errors"]+=1
        counts[status.lower()]+=1;output.append({"bib":str(get("bib")).strip(),"name":str(get("name")).strip() or "Sem nome","sex":str(get("sex")).strip(),"age":parse_int(get("age")),"category":str(get("category")).strip(),"distance_label":str(get("distance")).strip() or "Geral","gross_seconds":gross,"net_seconds":net,"overall_place":parse_int(get("overall_place")),"sex_place":parse_int(get("sex_place")),"category_place":parse_int(get("category_place")),"status":status,"raw_status":str(raw).strip()})
    counts["finishers"]=counts["finished"];return output,dict(counts)
def result_time(row):return row["net_seconds"] or row["gross_seconds"]
def percentile(values,p):
    if not values:return None
    values=sorted(values);k=(len(values)-1)*p;lo,hi=math.floor(k),math.ceil(k);return values[lo] if lo==hi else values[lo]*(hi-k)+values[hi]*(k-lo)
def summary(rows):
    times=[result_time(r) for r in rows if r["status"]=="FINISHED" and result_time(r)]
    return {"count":len(times),"best":min(times) if times else None,"mean":statistics.fmean(times) if times else None,"median":statistics.median(times) if times else None,"p10":percentile(times,.1),"p25":percentile(times,.25),"p50":percentile(times,.5),"p75":percentile(times,.75),"p90":percentile(times,.9),"times":times}
def smart_histogram(times):
    if not times:return []
    low,high=min(times),max(times);width=max(300,math.ceil(max(1,(high-low)/7)/300)*300);edge=math.floor(low/width)*width;bins=[]
    while edge<=high:
        end=edge+width;bins.append({"label":f"{format_seconds(edge)}–{format_seconds(end)}","count":sum(edge<=t<end for t in times)});edge=end
    maximum=max((b["count"] for b in bins),default=1)
    for b in bins:b["percent"]=round(b["count"]/maximum*100,1) if maximum else 0
    return bins
def calculate_stats(rows,distance_km=None):
    overall=summary(rows);sg,cg=defaultdict(list),defaultdict(list)
    for row in rows:
        if row["status"]=="FINISHED" and result_time(row):
            if row["sex"]:sg[row["sex"]].append(row)
            if row["category"]:cg[row["category"]].append(row)
    by_sex={k:summary(v) for k,v in sg.items()};by_category={k:summary(v) for k,v in cg.items()}
    for item in by_sex.values():item["share"]=round(item["count"]/overall["count"]*100,1) if overall["count"] else 0
    return {"overall":overall,"by_sex":by_sex,"by_category":by_category,"histogram":smart_histogram(overall["times"]),"statuses":Counter(r["status"] for r in rows),"distance_km":distance_km}
def athlete_analysis(athlete,rows):
    finished=sorted([r for r in rows if r["status"]=="FINISHED" and result_time(r)],key=result_time);t=result_time(athlete)
    def info(group,explicit):
        group=sorted(group,key=result_time);place=explicit or next((i+1 for i,r in enumerate(group) if r["id"]==athlete["id"]),None);total=len(group);return {"place":place,"total":total,"top":round(place/total*100,1) if place and total else None}
    sex=[r for r in finished if athlete["sex"] and r["sex"]==athlete["sex"]];cat=[r for r in finished if athlete["category"] and r["category"]==athlete["category"]];o,c=summary(finished),summary(cat);pace=t/athlete["distance_km"] if t and athlete["distance_km"] else None;compare=[{"label":"Seu tempo","value":t},{"label":"Mediana geral","value":o["median"]},{"label":"Média geral","value":o["mean"]},{"label":"Mediana da categoria","value":c["median"]},{"label":"Melhor da categoria","value":c["best"]}];compare=[x for x in compare if x["value"] is not None];mx=max((x["value"] for x in compare),default=1)
    for x in compare:x["percent"]=round(x["value"]/mx*100,1)
    return {"time":t,"pace":pace,"overall":info(finished,athlete["overall_place"]),"sex":info(sex,athlete["sex_place"]),"category":info(cat,athlete["category_place"]),"comparison":compare}
def import_path(token):
    if not re.fullmatch(r"[a-f0-9]{32}",token):abort(404)
    return Path(current_app.config["IMPORT_FOLDER"])/f"{token}.json"
def save_import(payload):
    token=uuid.uuid4().hex;write_import(token,payload);return token
def write_import(token,payload):import_path(token).write_text(json.dumps(payload,ensure_ascii=False),encoding="utf-8")
def load_import(token):
    path=import_path(token)
    if not path.exists():abort(404)
    return json.loads(path.read_text(encoding="utf-8"))
def delete_import(token):
    path=import_path(token)
    if path.exists():path.unlink()

SCHEMA="""
CREATE TABLE IF NOT EXISTS admins(id INTEGER PRIMARY KEY AUTOINCREMENT,username TEXT NOT NULL UNIQUE,password_hash TEXT NOT NULL,created_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS events(id INTEGER PRIMARY KEY AUTOINCREMENT,event_date TEXT NOT NULL,name TEXT NOT NULL,city TEXT NOT NULL,state TEXT NOT NULL CHECK(length(state)=2),registration_url TEXT NOT NULL DEFAULT '',photos_url TEXT NOT NULL DEFAULT '',status TEXT NOT NULL CHECK(status IN ('publicado','rascunho')),slug TEXT UNIQUE,created_at TEXT NOT NULL,updated_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS distances(id INTEGER PRIMARY KEY AUTOINCREMENT,event_id INTEGER NOT NULL,label TEXT NOT NULL,distance_km REAL,UNIQUE(event_id,label),FOREIGN KEY(event_id) REFERENCES events(id) ON DELETE CASCADE);
CREATE TABLE IF NOT EXISTS results(id INTEGER PRIMARY KEY AUTOINCREMENT,event_id INTEGER NOT NULL,distance_id INTEGER NOT NULL,bib TEXT,name TEXT NOT NULL,sex TEXT,age INTEGER,category TEXT,gross_seconds INTEGER,net_seconds INTEGER,overall_place INTEGER,sex_place INTEGER,category_place INTEGER,status TEXT NOT NULL DEFAULT 'FINISHED',raw_status TEXT,imported_at TEXT NOT NULL,FOREIGN KEY(event_id) REFERENCES events(id) ON DELETE CASCADE,FOREIGN KEY(distance_id) REFERENCES distances(id) ON DELETE CASCADE);
CREATE INDEX IF NOT EXISTS idx_events_public_date ON events(status,event_date);CREATE INDEX IF NOT EXISTS idx_results_distance_status ON results(distance_id,status);CREATE INDEX IF NOT EXISTS idx_results_search ON results(distance_id,name,bib);
"""
def ensure_schema():
    db=get_db();db.executescript(SCHEMA);columns={r[1] for r in db.execute("PRAGMA table_info(events)")}
    if "slug" not in columns:db.execute("ALTER TABLE events ADD COLUMN slug TEXT")
    for row in db.execute("SELECT id,name FROM events WHERE slug IS NULL OR slug='' ").fetchall():db.execute("UPDATE events SET slug=? WHERE id=?",(unique_slug(row["name"],row["id"]),row["id"]))
    db.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_events_slug ON events(slug)");db.commit()
def init_database(seed=True):
    ensure_schema();db=get_db();username=os.environ.get("ADMIN_USERNAME","admin");password=os.environ.get("ADMIN_PASSWORD")
    if not password:raise click.ClickException("Defina ADMIN_PASSWORD antes de inicializar o banco.")
    db.execute("INSERT OR IGNORE INTO admins(username,password_hash,created_at) VALUES(?,?,?)",(username,generate_password_hash(password),datetime.now().isoformat(timespec="seconds")))
    if seed and db.execute("SELECT COUNT(*) FROM events").fetchone()[0]==0:
        today=date.today();now=datetime.now().isoformat(timespec="seconds");samples=[(today+timedelta(days=3),"Joinville Night Run","Joinville","SC","https://example.com/inscricao","","publicado"),(today+timedelta(days=11),"Circuito das Águas","Blumenau","SC","https://example.com/circuito","","publicado"),(today-timedelta(days=4),"Desafio da Cidade","Curitiba","PR","","https://example.com/fotos","publicado"),(today+timedelta(days=25),"Meia Maratona do Litoral","Itajaí","SC","","","rascunho")]
        db.executemany("INSERT INTO events(event_date,name,city,state,registration_url,photos_url,status,slug,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?)",[(d.isoformat(),n,c,s,i,p,st,unique_slug(n),now,now) for d,n,c,s,i,p,st in samples])
    db.commit()
@click.command("init-db")
@click.option("--sem-exemplos",is_flag=True)
def init_db_command(sem_exemplos):init_database(seed=not sem_exemplos);click.echo("Banco inicializado com sucesso.")
app=create_app()
if __name__=="__main__":app.run(debug=os.environ.get("FLASK_DEBUG")=="1")
